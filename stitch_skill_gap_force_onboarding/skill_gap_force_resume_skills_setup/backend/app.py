"""
Skill Gap Force -- Resume Skill Extractor & Job Category Classifier Backend
=============================================================================
A lightweight Flask API that accepts a resume upload (PDF / DOC / DOCX),
extracts raw text from it, performs case-insensitive whole-word keyword
matching against a curated skills catalog, and classifies the profile into
one of three primary job categories:
  - web_dev ("Web Development")
  - ai_ml   ("AI & Machine Learning")
  - web3    ("Web3 & Blockchain")

No AI / LLM involved -- pure keyword matching with robust word-boundary regexes.
"""

import re
import io
import os

from flask import Flask, request, jsonify
from flask_cors import CORS

# Optional parsers
try:
    import pdfplumber
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# ---------------------------------------------------------------------------
# Flask App Setup
# ---------------------------------------------------------------------------
app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "*"}}, supports_credentials=True)

@app.after_request
def add_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS, PUT, DELETE"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, X-Requested-With, Origin, Accept"
    return response

# ---------------------------------------------------------------------------
# Category Keyword Lists (Step 1)
# ---------------------------------------------------------------------------
WEB_DEV_SKILLS = [
    "HTML", "CSS", "JavaScript", "Git", "GitHub", "React", "Node.js", "Express",
    "Database", "REST API", "Authentication", "TypeScript", "Next.js", "Docker",
    "Cloud", "System Design"
]

AI_ML_SKILLS = [
    "Python", "NumPy", "Pandas", "Matplotlib", "Seaborn", "Statistics", "Probability",
    "Linear Algebra", "SQL", "Data Preprocessing", "EDA", "Machine Learning",
    "Scikit-learn", "Supervised Learning", "Unsupervised Learning", "Feature Engineering",
    "Model Evaluation", "Deep Learning", "PyTorch", "TensorFlow", "CNN", "RNN", "LSTM",
    "Transformers", "NLP", "GenAI", "LLM", "Hugging Face", "LangChain", "RAG",
    "Fine-tuning", "MLOps", "Docker", "AWS", "GCP", "Azure"
]

WEB3_SKILLS = [
    "Blockchain", "Cryptography", "Bitcoin", "Ethereum", "Wallets", "EVM", "Solidity",
    "Smart Contracts", "Remix", "Hardhat", "Foundry", "Web3.js", "Ethers.js", "dApps",
    "ERC-20", "ERC-721", "ERC-1155", "DeFi", "DEX", "DAOs", "Oracles", "Layer 2",
    "Gas Optimization", "Smart Contract Security", "Testing", "Auditing", "IPFS",
    "The Graph", "RPC", "MEV", "Tokenomics", "Cross-chain", "Web3 Backend", "Docker",
    "Cloud"
]

CATEGORY_NAMES = {
    "web_dev": "Web Development",
    "ai_ml": "AI & Machine Learning",
    "web3": "Web3 & Blockchain"
}

# ---------------------------------------------------------------------------
# Skills Catalog (Canonical Name, [Aliases / Variants])
# Ensures all category keywords + industry skills are extractable
# ---------------------------------------------------------------------------
SKILLS_CATALOG = [
    # ── Web Development Keywords ──
    ("HTML", ["HTML", "HTML5"]),
    ("CSS", ["CSS", "CSS3"]),
    ("JavaScript", ["JavaScript", "JS"]),
    ("TypeScript", ["TypeScript", "TS"]),
    ("React", ["React", "React.js", "ReactJS"]),
    ("Node.js", ["Node.js", "NodeJS", "Node"]),
    ("Express", ["Express", "Express.js", "ExpressJS"]),
    ("Next.js", ["Next.js", "NextJS", "Next"]),
    ("Database", ["Database", "Databases", "DBMS", "RDBMS"]),
    ("REST API", ["REST API", "REST APIs", "RESTful API", "RESTful APIs", "RESTful", "REST"]),
    ("Authentication", ["Authentication", "Auth", "OAuth", "JWT"]),
    ("System Design", ["System Design"]),
    ("Cloud", ["Cloud", "Cloud Computing"]),
    ("Git", ["Git"]),
    ("GitHub", ["GitHub"]),
    ("Docker", ["Docker"]),

    # ── AI / ML Keywords ──
    ("Python", ["Python"]),
    ("NumPy", ["NumPy"]),
    ("Pandas", ["Pandas"]),
    ("Matplotlib", ["Matplotlib"]),
    ("Seaborn", ["Seaborn"]),
    ("Statistics", ["Statistics", "Statistical Modeling", "Statistical Analysis"]),
    ("Probability", ["Probability"]),
    ("Linear Algebra", ["Linear Algebra"]),
    ("SQL", ["SQL"]),
    ("Data Preprocessing", ["Data Preprocessing", "Data Cleaning"]),
    ("EDA", ["EDA", "Exploratory Data Analysis"]),
    ("Machine Learning", ["Machine Learning", "ML"]),
    ("Scikit-learn", ["Scikit-learn", "Scikit learn", "Sklearn"]),
    ("Supervised Learning", ["Supervised Learning"]),
    ("Unsupervised Learning", ["Unsupervised Learning"]),
    ("Feature Engineering", ["Feature Engineering"]),
    ("Model Evaluation", ["Model Evaluation"]),
    ("Deep Learning", ["Deep Learning", "DL"]),
    ("PyTorch", ["PyTorch"]),
    ("TensorFlow", ["TensorFlow"]),
    ("CNN", ["CNN", "Convolutional Neural Networks", "Convolutional Neural Network"]),
    ("RNN", ["RNN", "Recurrent Neural Networks", "Recurrent Neural Network"]),
    ("LSTM", ["LSTM", "Long Short-Term Memory"]),
    ("Transformers", ["Transformers", "Transformer Models", "Transformer"]),
    ("NLP", ["NLP", "Natural Language Processing"]),
    ("GenAI", ["GenAI", "Generative AI"]),
    ("LLM", ["LLM", "LLMs", "Large Language Models", "Large Language Model"]),
    ("Hugging Face", ["Hugging Face", "HuggingFace"]),
    ("LangChain", ["LangChain"]),
    ("RAG", ["RAG", "Retrieval-Augmented Generation", "Retrieval Augmented Generation"]),
    ("Fine-tuning", ["Fine-tuning", "Fine tuning", "Model Fine-Tuning"]),
    ("MLOps", ["MLOps", "Machine Learning Operations"]),
    ("AWS", ["AWS", "Amazon Web Services"]),
    ("GCP", ["GCP", "Google Cloud Platform", "Google Cloud"]),
    ("Azure", ["Azure", "Microsoft Azure"]),

    # ── Web3 Keywords ──
    ("Blockchain", ["Blockchain", "Blockchains"]),
    ("Cryptography", ["Cryptography", "Cryptographic"]),
    ("Bitcoin", ["Bitcoin", "BTC"]),
    ("Ethereum", ["Ethereum", "ETH"]),
    ("Wallets", ["Wallets", "Crypto Wallets", "Metamask"]),
    ("EVM", ["EVM", "Ethereum Virtual Machine"]),
    ("Solidity", ["Solidity"]),
    ("Smart Contracts", ["Smart Contracts", "Smart Contract"]),
    ("Remix", ["Remix", "Remix IDE"]),
    ("Hardhat", ["Hardhat"]),
    ("Foundry", ["Foundry"]),
    ("Web3.js", ["Web3.js", "Web3js", "Web3"]),
    ("Ethers.js", ["Ethers.js", "Ethersjs", "Ethers"]),
    ("dApps", ["dApps", "dApp", "Decentralized Applications", "Decentralized App"]),
    ("ERC-20", ["ERC-20", "ERC20"]),
    ("ERC-721", ["ERC-721", "ERC721"]),
    ("ERC-1155", ["ERC-1155", "ERC1155"]),
    ("DeFi", ["DeFi", "Decentralized Finance"]),
    ("DEX", ["DEX", "Decentralized Exchanges", "Decentralized Exchange"]),
    ("DAOs", ["DAOs", "DAO", "Decentralized Autonomous Organizations"]),
    ("Oracles", ["Oracles", "Oracle", "Chainlink Oracles"]),
    ("Layer 2", ["Layer 2", "L2", "Rollups", "Arbitrum", "Optimism", "Polygon"]),
    ("Gas Optimization", ["Gas Optimization", "Gas Efficient"]),
    ("Smart Contract Security", ["Smart Contract Security", "Smart Contract Audit"]),
    ("Testing", ["Testing", "Unit Testing", "Integration Testing"]),
    ("Auditing", ["Auditing", "Security Auditing", "Code Audit"]),
    ("IPFS", ["IPFS", "InterPlanetary File System"]),
    ("The Graph", ["The Graph", "Subgraphs", "Subgraph"]),
    ("RPC", ["RPC", "JSON-RPC"]),
    ("MEV", ["MEV", "Maximal Extractable Value"]),
    ("Tokenomics", ["Tokenomics"]),
    ("Cross-chain", ["Cross-chain", "Cross chain", "Bridges"]),
    ("Web3 Backend", ["Web3 Backend"]),

    # ── Other Popular Tech, Tools & Soft Skills ──
    ("Java", ["Java"]),
    ("C++", ["C++", "Cpp"]),
    ("C#", ["C#", "CSharp", "C-Sharp"]),
    ("C", ["C"]),
    ("Go", ["Golang", "Go"]),
    ("Rust", ["Rust"]),
    ("Swift", ["Swift"]),
    ("Kotlin", ["Kotlin"]),
    ("Ruby", ["Ruby"]),
    ("PHP", ["PHP"]),
    ("Tailwind CSS", ["Tailwind CSS", "TailwindCSS", "Tailwind"]),
    ("Vue.js", ["Vue.js", "VueJS", "Vue"]),
    ("Angular", ["Angular", "AngularJS"]),
    ("GraphQL", ["GraphQL"]),
    ("Kubernetes", ["Kubernetes", "K8s"]),
    ("Linux", ["Linux"]),
    ("Figma", ["Figma"]),
    ("Excel", ["Excel", "MS Excel"]),
    ("Communication", ["Communication"]),
    ("Leadership", ["Leadership"]),
    ("Project Management", ["Project Management"]),
    ("Agile", ["Agile"]),
    ("Scrum", ["Scrum"]),
    ("Problem Solving", ["Problem Solving"]),
    ("Data Analysis", ["Data Analysis"]),
    ("Data Science", ["Data Science"]),
]


# ---------------------------------------------------------------------------
# Regex Compilation with Precise Word Boundary Rules
# ---------------------------------------------------------------------------
def _compile_pattern(term: str) -> re.Pattern:
    escaped = re.escape(term)
    if term.endswith("++") or term.endswith("#"):
        return re.compile(r'(?<![a-zA-Z0-9])' + escaped + r'(?![a-zA-Z0-9+#])', re.IGNORECASE)
    elif term == "C":
        return re.compile(r'(?<![a-zA-Z0-9])C(?![a-zA-Z0-9+#])', re.IGNORECASE)
    elif term == "R":
        return re.compile(r'(?<![a-zA-Z0-9])R(?![a-zA-Z0-9])', re.IGNORECASE)
    else:
        return re.compile(r'(?<![a-zA-Z0-9])' + escaped + r'(?![a-zA-Z0-9])', re.IGNORECASE)


COMPILED_SKILLS = []
for canonical_name, variants in SKILLS_CATALOG:
    patterns = [_compile_pattern(v) for v in variants]
    COMPILED_SKILLS.append((canonical_name, patterns))


# ---------------------------------------------------------------------------
# Text Extraction Helpers
# ---------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    if not HAS_PDF:
        raise RuntimeError("pdfplumber is not installed. Please run: pip install pdfplumber")
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_bytes: bytes) -> str:
    if not HAS_DOCX:
        raise RuntimeError("python-docx is not installed. Please run: pip install python-docx")
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception:
        text_tokens = re.findall(rb'[\x20-\x7E]{3,}', file_bytes)
        if text_tokens:
            return " ".join([t.decode('latin1', errors='ignore') for t in text_tokens])
        raise


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext in ("doc", "docx"):
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file format '.{ext}'. Supported: PDF, DOC, DOCX.")


# ---------------------------------------------------------------------------
# Skill Matching & Category Classification (Steps 2 & 3)
# ---------------------------------------------------------------------------
def match_skills(raw_text: str) -> list:
    matched = []
    seen = set()
    for canonical_name, patterns in COMPILED_SKILLS:
        for pattern in patterns:
            if pattern.search(raw_text):
                key = canonical_name.lower()
                if key not in seen:
                    seen.add(key)
                    matched.append(canonical_name)
                break
    return matched


def _norm_token(s: str) -> str:
    return s.lower().replace(".", "").replace("-", " ").replace("_", " ").strip()


def skill_matches_keyword(skill: str, keyword: str) -> bool:
    s_norm = _norm_token(skill)
    k_norm = _norm_token(keyword)
    if s_norm == k_norm:
        return True
    if s_norm + "s" == k_norm or k_norm + "s" == s_norm:
        return True
    return False


def classify_skills(skills: list) -> dict:
    """
    Counts matches against WEB_DEV_SKILLS, AI_ML_SKILLS, WEB3_SKILLS.
    Determines single winning category:
      - Highest match count wins
      - If tie: highest percentage of category's list size wins
      - If all 0: returns None (edge case fallback)
    """
    web_dev_matches = [s for s in skills if any(skill_matches_keyword(s, kw) for kw in WEB_DEV_SKILLS)]
    ai_ml_matches = [s for s in skills if any(skill_matches_keyword(s, kw) for kw in AI_ML_SKILLS)]
    web3_matches = [s for s in skills if any(skill_matches_keyword(s, kw) for kw in WEB3_SKILLS)]

    counts = {
        "web_dev": len(web_dev_matches),
        "ai_ml": len(ai_ml_matches),
        "web3": len(web3_matches)
    }
    lengths = {
        "web_dev": len(WEB_DEV_SKILLS),
        "ai_ml": len(AI_ML_SKILLS),
        "web3": len(WEB3_SKILLS)
    }
    percentages = {
        k: round((counts[k] / lengths[k]) * 100, 1) for k in counts
    }

    max_count = max(counts.values())

    if max_count == 0:
        winning_category = None
        winning_name = None
    else:
        candidates = [k for k, v in counts.items() if v == max_count]
        if len(candidates) == 1:
            winning_category = candidates[0]
        else:
            # Exact tie breaker: percentage of that category's list size
            winner_candidate = max(candidates, key=lambda c: counts[c] / lengths[c])
            winning_category = winner_candidate
        winning_name = CATEGORY_NAMES.get(winning_category)

    return {
        "winning_category": winning_category,
        "winning_category_name": winning_name,
        "web_dev_match_count": counts["web_dev"],
        "ai_ml_match_count": counts["ai_ml"],
        "web3_match_count": counts["web3"],
        "web_dev_percentage": percentages["web_dev"],
        "ai_ml_percentage": percentages["ai_ml"],
        "web3_percentage": percentages["web3"],
        "matched_skills": {
            "web_dev": web_dev_matches,
            "ai_ml": ai_ml_matches,
            "web3": web3_matches
        }
    }


# ---------------------------------------------------------------------------
# Additional Profile Details Extraction (Phone, Email, Location, Education, Experience)
# ---------------------------------------------------------------------------
def extract_email(text: str) -> str:
    match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    if match:
        email = match.group(0).strip().rstrip('.')
        if not email.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp')):
            return email
    return ""


def extract_phone(text: str) -> str:
    phone_label_match = re.search(r'(?:phone|tel|mobile|cell|contact)\s*[:\-]?\s*([+\d\s().-]{10,25})', text, re.IGNORECASE)
    if phone_label_match:
        cand = phone_label_match.group(1).strip()
        digits = re.sub(r'\D', '', cand)
        if 10 <= len(digits) <= 15:
            return cand

    patterns = [
        r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',
        r'\+?\d{1,4}[-.\s]?\d{4,5}[-.\s]?\d{4,5}',
        r'\b\d{10}\b'
    ]
    for p in patterns:
        for m in re.finditer(p, text):
            cand = m.group(0).strip()
            digits = re.sub(r'\D', '', cand)
            if 10 <= len(digits) <= 15:
                return cand
    return ""


def extract_location(text: str) -> str:
    label_match = re.search(r'(?:location|address|city|residing in|based in)\s*[:\-]\s*([^\n\r,]+(?:,\s*[^\n\r]+)?)', text, re.IGNORECASE)
    if label_match:
        loc = label_match.group(1).strip()
        if 2 < len(loc) < 60:
            return loc

    top_text = "\n".join(text.splitlines()[:30])
    common_locations = [
        r'San Francisco(?:,\s*CA)?', r'San Jose(?:,\s*CA)?', r'Silicon Valley(?:,\s*CA)?',
        r'New York(?:,\s*NY)?', r'Seattle(?:,\s*WA)?', r'Austin(?:,\s*TX)?',
        r'Boston(?:,\s*MA)?', r'Chicago(?:,\s*IL)?', r'Los Angeles(?:,\s*CA)?',
        r'London(?:,\s*UK)?', r'Toronto(?:,\s*Canada)?', r'Vancouver(?:,\s*Canada)?',
        r'Bengaluru(?:,\s*India)?', r'Bangalore(?:,\s*India)?', r'Mumbai(?:,\s*India)?',
        r'Delhi(?:,\s*India)?', r'Hyderabad(?:,\s*India)?', r'Pune(?:,\s*India)?',
        r'Remote', r'Berlin(?:,\s*Germany)?', r'Singapore'
    ]
    for loc_pat in common_locations:
        m = re.search(r'\b' + loc_pat + r'\b', top_text, re.IGNORECASE)
        if m:
            return m.group(0).strip()
    
    for line in top_text.splitlines():
        line = line.strip()
        city_state_match = re.search(r'\b([A-Z][a-zA-Z\s]{2,20},\s*(?:[A-Z]{2}|[A-Z][a-zA-Z]{3,15}))\b', line)
        if city_state_match and not any(kw in line.lower() for kw in ['university', 'college', 'school', 'skills', 'experience']):
            return city_state_match.group(1).strip()

    return ""


def extract_education(text: str) -> str:
    degree_patterns = [
        r'(?:B\.?Tech|B\.?E\.?|B\.?S\.?|B\.?Sc|Bachelor(?:\'s)?(?:\s+of|\s+in)?\s+[A-Za-z\s]+)(?:[^\n\r;,]{0,60})',
        r'(?:M\.?Tech|M\.?E\.?|M\.?S\.?|M\.?Sc|Master(?:\'s)?(?:\s+of|\s+in)?\s+[A-Za-z\s]+)(?:[^\n\r;,]{0,60})',
        r'(?:M\.?B\.?A\.?|MBA)(?:[^\n\r;,]{0,40})',
        r'(?:Ph\.?D\.?|Doctorate)(?:[^\n\r;,]{0,50})',
        r'(?:Diploma|Associate(?:\'s)?\s+Degree)(?:[^\n\r;,]{0,50})'
    ]
    for p in degree_patterns:
        match = re.search(p, text, re.IGNORECASE)
        if match:
            edu = match.group(0).strip()
            # Stop if other section labels are caught in the degree tail
            edu = re.split(r'\b(?:Experience|Skills|Projects|Work|Employment|Phone|Email|Location)\s*[:\-]', edu, flags=re.IGNORECASE)[0].strip()
            edu = re.sub(r'\s+', ' ', edu).rstrip(';,.- ')
            if 3 < len(edu) < 100:
                return edu

    lines = text.splitlines()
    in_edu = False
    for line in lines:
        clean_line = line.strip()
        if re.match(r'^(?:Education|Academic|Qualification|Degrees)\b', clean_line, re.IGNORECASE):
            in_edu = True
            continue
        if in_edu and clean_line:
            if re.match(r'^(?:Experience|Projects|Skills|Certifications|Work|Summary)\b', clean_line, re.IGNORECASE):
                break
            if any(kw in clean_line.lower() for kw in ['university', 'college', 'institute', 'bachelor', 'master', 'degree', 'b.tech', 'b.e', 'b.s']):
                return clean_line[:80]
    return ""


def extract_experience(text: str) -> str:
    exp_years_match = re.search(r'\b(\d+[\+]?\s*(?:years|yrs)(?:\s+of)?\s+(?:professional\s+)?(?:experience|exp))\b', text, re.IGNORECASE)
    if exp_years_match:
        return exp_years_match.group(1).strip()

    titles = [
        "Senior Full Stack Engineer", "Senior Full Stack Developer", "Senior Frontend Engineer",
        "Senior Backend Engineer", "Senior Software Engineer", "Full Stack Engineer",
        "Full Stack Developer", "Frontend Engineer", "Frontend Developer",
        "Backend Engineer", "Backend Developer", "Software Engineer",
        "Machine Learning Engineer", "AI Engineer", "Data Scientist",
        "Data Engineer", "DevOps Engineer", "Cloud Architect",
        "Blockchain Developer", "Smart Contract Engineer", "Web3 Developer",
        "Mobile Developer", "iOS Developer", "Android Developer",
        "Product Engineer", "Engineering Manager"
    ]
    top_lines = text.splitlines()[:25]
    for line in top_lines:
        clean = line.strip()
        for title in titles:
            if re.search(r'\b' + re.escape(title) + r'\b', clean, re.IGNORECASE):
                return title

    exp_label_match = re.search(r'(?:experience|years of experience|total exp)\s*[:\-]\s*([^\n\r]+)', text, re.IGNORECASE)
    if exp_label_match:
        val = exp_label_match.group(1).strip()
        if 2 < len(val) < 40:
            return val

    return ""


def extract_profile_details(text: str) -> dict:
    return {
        "phone": extract_phone(text),
        "email": extract_email(text),
        "location": extract_location(text),
        "education": extract_education(text),
        "experience": extract_experience(text)
    }


# ---------------------------------------------------------------------------
# API Endpoints
# ---------------------------------------------------------------------------
@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "status": "online",
        "total_skills_supported": len(COMPILED_SKILLS),
        "pdfplumber": HAS_PDF,
        "python_docx": HAS_DOCX
    }), 200


@app.route("/extract", methods=["POST", "OPTIONS"])
@app.route("/extract-skills", methods=["POST", "OPTIONS"])
def extract_endpoint():
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"}), 200

    # Accept field name 'file' (required) or 'resume' (fallback)
    file_obj = request.files.get("file") or request.files.get("resume")
    if not file_obj:
        print("[Upload Error]: No file uploaded. Expected field name 'file'.")
        return jsonify({"error": "No file uploaded. Please upload a resume using field name 'file'."}), 400

    filename = file_obj.filename or ""
    if not filename:
        print("[Upload Error]: Uploaded file is missing a filename.")
        return jsonify({"error": "Uploaded file is missing a filename."}), 400

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ("pdf", "doc", "docx"):
        print(f"[Upload Error]: Unsupported file format '.{ext}' for file '{filename}'.")
        return jsonify({"error": f"Unsupported file format '.{ext}'. Supported formats: PDF, DOC, DOCX."}), 400

    file_bytes = file_obj.read()
    if len(file_bytes) == 0:
        print(f"[Upload Error]: Uploaded file '{filename}' is empty (0 bytes).")
        return jsonify({"error": "The uploaded file is empty (0 bytes)."}), 400

    if len(file_bytes) > 5 * 1024 * 1024:
        print(f"[Upload Error]: Uploaded file '{filename}' exceeds 5MB limit ({len(file_bytes)} bytes).")
        return jsonify({"error": "File size exceeds 5MB limit."}), 400

    try:
        extracted_text = extract_text(file_bytes, filename)
    except RuntimeError as e:
        import traceback
        traceback.print_exc()
        print(f"[Extraction Error]: {str(e)}")
        return jsonify({"error": str(e)}), 500
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"[Extraction Error]: Failed to extract text: {str(e)}")
        return jsonify({"error": f"Failed to extract text: {str(e)}"}), 500

    if not extracted_text or not extracted_text.strip():
        empty_classification = classify_skills([])
        print(f"[Extraction Warning]: No readable text extracted from '{filename}'.")
        return jsonify({
            "skills": [],
            "count": 0,
            "filename": filename,
            "category_classification": empty_classification,
            "phone": "",
            "email": "",
            "location": "",
            "education": "",
            "experience": "",
            "message": "No readable text found in document."
        }), 200

    skills = match_skills(extracted_text)
    classification = classify_skills(skills)
    profile_details = extract_profile_details(extracted_text)

    print(f"[Extraction Success]: Successfully extracted {len(skills)} skills from '{filename}'.")
    return jsonify({
        "skills": skills,
        "count": len(skills),
        "filename": filename,
        "category_classification": classification,
        "phone": profile_details.get("phone") or "",
        "email": profile_details.get("email") or "",
        "location": profile_details.get("location") or "",
        "education": profile_details.get("education") or "",
        "experience": profile_details.get("experience") or ""
    }), 200


# ---------------------------------------------------------------------------
# Server Entrypoint
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    print("=" * 64)
    print("  Skill Gap Force -- Resume Extractor & Classifier")
    print(f"  Running on: http://localhost:{port}")
    print(f"  Supported Skills: {len(COMPILED_SKILLS)}")
    print(f"  pdfplumber : {'Available' if HAS_PDF else 'Missing'}")
    print(f"  python-docx: {'Available' if HAS_DOCX else 'Missing'}")
    print("=" * 64)
    app.run(host="0.0.0.0", port=port, debug=False)
